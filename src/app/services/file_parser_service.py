import csv
import io
from typing import Any
from fastapi import HTTPException, UploadFile


REQUIRED_LEAD_COLUMNS = [
    "name",
    "email",
    "phone",
    "source",
    "property_type",
    "configuration",
    "location_preference",
    "budget",
    "timeline",
    "message",
]

OPTIONAL_LEAD_COLUMNS = [
    "call_consent",
    "do_not_call",
]


class FileParserService:
    @staticmethod
    async def parse_lead_upload(file: UploadFile) -> list[dict[str, Any]]:
        if not file.filename:
            raise HTTPException(status_code=400, detail="File name is required.")

        filename = file.filename.lower()
        content = await file.read()

        if filename.endswith(".csv"):
            rows = FileParserService._parse_csv(content)
        elif filename.endswith(".xlsx"):
            rows = FileParserService._parse_xlsx(content)
        else:
            raise HTTPException(status_code=400, detail="Only .csv and .xlsx files are supported.")

        normalized_rows = [
            {str(key).strip().lower(): value for key, value in row.items()}
            for row in rows
        ]
        FileParserService._validate_columns(normalized_rows)
        return [
            {key: str(row.get(key, "") or "").strip() for key in [*REQUIRED_LEAD_COLUMNS, *OPTIONAL_LEAD_COLUMNS]}
            for row in normalized_rows
        ]

    @staticmethod
    def _parse_csv(content: bytes) -> list[dict[str, Any]]:
        try:
            decoded = content.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise HTTPException(status_code=400, detail="CSV must be UTF-8 encoded.") from exc

        reader = csv.DictReader(io.StringIO(decoded))
        return list(reader)

    @staticmethod
    def _parse_xlsx(content: bytes) -> list[dict[str, Any]]:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise HTTPException(status_code=500, detail="Excel support requires openpyxl.") from exc

        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return []

        headers = [str(cell or "").strip() for cell in rows[0]]
        parsed = []
        for values in rows[1:]:
            parsed.append({headers[index]: value for index, value in enumerate(values) if index < len(headers)})
        return parsed

    @staticmethod
    def _validate_columns(rows: list[dict[str, Any]]) -> None:
        if not rows:
            raise HTTPException(status_code=400, detail="Uploaded file has no lead rows.")

        normalized_headers = {str(header).strip().lower() for header in rows[0].keys()}
        missing = [column for column in REQUIRED_LEAD_COLUMNS if column not in normalized_headers]
        if missing:
            raise HTTPException(
                status_code=400,
                detail=f"Missing required columns: {', '.join(missing)}",
            )
